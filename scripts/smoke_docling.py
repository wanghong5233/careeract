import argparse
import os
import resource
import time
from pathlib import Path


def create_pdf(path: Path) -> None:
    stream = b"BT /F1 18 Tf 72 720 Td (CareerAct synthetic resume) Tj ET"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    content = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, value in enumerate(objects, start=1):
        offsets.append(len(content))
        content.extend(f"{index} 0 obj\n".encode() + value + b"\nendobj\n")
    xref = len(content)
    content.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        content.extend(f"{offset:010d} 00000 n \n".encode())
    content.extend(
        f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    path.write_bytes(content)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--directory", type=Path, default=Path("/data"))
    parser.add_argument("--offline", action="store_true")
    args = parser.parse_args()
    if args.offline:
        os.environ["HF_HUB_OFFLINE"] = "1"
    os.environ.setdefault("HF_HOME", str(args.directory / "models"))

    from docling.datamodel.accelerator_options import AcceleratorOptions
    from docling.datamodel.base_models import ConversionStatus, InputFormat
    from docling.datamodel.pipeline_options import PdfPipelineOptions
    from docling.document_converter import DocumentConverter, PdfFormatOption
    from docx import Document

    args.directory.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_heading("CareerAct synthetic resume", level=1)
    document.add_paragraph("虚构候选人：测试同学。技能：Python、人工智能。")
    docx_path = args.directory / "synthetic.docx"
    document.save(str(docx_path))
    pdf_path = args.directory / "synthetic.pdf"
    create_pdf(pdf_path)
    options = PdfPipelineOptions(
        do_ocr=False,
        do_table_structure=False,
        document_timeout=180,
        accelerator_options=AcceleratorOptions(device="cpu", num_threads=2),
    )
    converter = DocumentConverter(
        allowed_formats=[InputFormat.DOCX, InputFormat.PDF],
        format_options={InputFormat.PDF: PdfFormatOption(pipeline_options=options)},
    )
    for path, expected in [(docx_path, "测试同学"), (pdf_path, "CareerAct synthetic resume")]:
        started = time.monotonic()
        result = converter.convert(path)
        if result.status != ConversionStatus.SUCCESS:
            raise RuntimeError(f"{path.suffix} conversion status: {result.status}")
        if expected not in result.document.export_to_markdown():
            raise RuntimeError(f"{path.suffix} expected synthetic text missing")
        print(
            f"PASS: {path.suffix} content verified in {time.monotonic() - started:.1f}s", flush=True
        )
    print(f"PASS: offline={args.offline}; OCR and table recognition not tested", flush=True)
    print(f"Process peak RSS: {resource.getrusage(resource.RUSAGE_SELF).ru_maxrss / 1024:.1f} MiB")


if __name__ == "__main__":
    main()
